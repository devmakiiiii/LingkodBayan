"""Build LingkodBayan.pptx from the presentation content (mirrors slides.html)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import re, os

FONT = 'Segoe UI'
BG = RGBColor(0x0F, 0x17, 0x2A)
CARD = RGBColor(0x1E, 0x29, 0x3B)
BORDER = RGBColor(0x33, 0x41, 0x55)
ACCENT = RGBColor(0x38, 0xBD, 0xF8)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
TEXT = RGBColor(0xE2, 0xE8, 0xF0)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
AMBER = RGBColor(0xFB, 0xBF, 0x24)
PILL_BLUE = RGBColor(0x0C, 0x4A, 0x6E)
PILL_GREEN = RGBColor(0x14, 0x53, 0x2D)
CYAN = RGBColor(0x67, 0xE8, 0xF9)
TAG_BG = RGBColor(0x16, 0x4E, 0x63)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 14


def new_slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def textbox(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def rich(p, text, size, color=TEXT, bold=False, font=FONT):
    """Add runs to paragraph p, honoring **bold** segments (bold -> ACCENT)."""
    for idx, part in enumerate(text.split('**')):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        r.font.size = Pt(size)
        r.font.name = font
        is_bold = bold or (idx % 2 == 1)
        r.font.bold = is_bold
        r.font.color.rgb = ACCENT if idx % 2 == 1 else color


def bullet(tf, text, size=17, first=False, color=TEXT, space_after=10, indent_marker='• '):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    rich(p, indent_marker + text, size, color=color)
    return p


def add_title(slide, num, text):
    tf = textbox(slide, 0.7, 0.4, 12.2, 1.0)
    p = tf.paragraphs[0]
    if num:
        r = p.add_run()
        r.text = num + '  '
        r.font.size = Pt(30)
        r.font.name = FONT
        r.font.color.rgb = MUTED
    r = p.add_run()
    r.text = text
    r.font.size = Pt(34)
    r.font.name = FONT
    r.font.bold = True
    r.font.color.rgb = ACCENT


def footer(slide, n):
    tf = textbox(slide, 0.7, 6.95, 8.0, 0.4)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = 'LingkodBayan — Civic Services Portal'
    r.font.size = Pt(10)
    r.font.name = FONT
    r.font.color.rgb = MUTED
    tf2 = textbox(slide, 11.6, 6.95, 1.3, 0.4)
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = f'{n} / {TOTAL}'
    r2.font.size = Pt(10)
    r2.font.name = FONT
    r2.font.color.rgb = MUTED


def no_shadow(shape):
    shape.shadow.inherit = False


def card(slide, l, t, w, h, title, items, title_color=GREEN, size=15):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = CARD
    box.line.color.rgb = BORDER
    box.line.width = Pt(1)
    no_shadow(box)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.28)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.22)
    p = tf.paragraphs[0]
    p.space_after = Pt(10)
    r = p.add_run()
    r.text = title
    r.font.size = Pt(size + 3)
    r.font.name = FONT
    r.font.bold = True
    r.font.color.rgb = title_color
    for it in items:
        bullet(tf, it, size=size, space_after=7)
    return box


def pill(slide, l, t, w, h, text, terminal=False, size=14):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    try:
        sh.adjustments[0] = 0.5
    except Exception:
        pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = PILL_GREEN if terminal else PILL_BLUE
    sh.line.color.rgb = GREEN if terminal else ACCENT
    sh.line.width = Pt(1.5)
    no_shadow(sh)
    tf = sh.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = FONT
    r.font.bold = True
    r.font.color.rgb = TEXT
    return sh


def arrow(slide, l, t, w=0.5, h=0.6, txt='→', size=22, color=ACCENT):
    tf = textbox(slide, l, t, w, h)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = txt
    r.font.size = Pt(size)
    r.font.name = FONT
    r.font.color.rgb = color


# ---------- Slide 1: Title ----------
s = new_slide()
tf = textbox(s, 0.9, 2.0, 11.5, 1.6)
p = tf.paragraphs[0]
r = p.add_run(); r.text = 'LingkodBayan'
r.font.size = Pt(60); r.font.name = FONT; r.font.bold = True; r.font.color.rgb = ACCENT
tf = textbox(s, 0.9, 3.4, 11.5, 0.9)
p = tf.paragraphs[0]
r = p.add_run(); r.text = 'A Civic Services Portal for Connecting Citizens with Local Government'
r.font.size = Pt(24); r.font.name = FONT; r.font.color.rgb = MUTED
tf = textbox(s, 0.9, 4.5, 11.5, 0.6)
p = tf.paragraphs[0]
r = p.add_run(); r.text = 'Next.js 16   ·   React 19   ·   TypeScript   ·   Supabase   ·   Vercel'
r.font.size = Pt(18); r.font.name = FONT; r.font.color.rgb = CYAN
tf = textbox(s, 0.9, 5.6, 11.5, 0.5)
p = tf.paragraphs[0]
r = p.add_run(); r.text = 'Student Presentation · github.com/devmakiiiii/LingkodBayan'
r.font.size = Pt(14); r.font.name = FONT; r.font.color.rgb = MUTED
footer(s, 1)

# ---------- Slide 2: The Problem ----------
s = new_slide()
add_title(s, '01', 'The Problem')
tf = textbox(s, 0.9, 1.7, 11.5, 4.2)
items = [
    'Citizens must **queue in person** for simple services — water, waste, street repairs',
    'Complaints have **no transparent tracking** — "did anyone even read it?"',
    'Announcements reach residents **inconsistently**',
    'Admins manage everything with **paper and spreadsheets**',
]
for i, it in enumerate(items):
    bullet(tf, it, size=20, first=(i == 0), space_after=16)
tf = textbox(s, 0.9, 5.7, 11.5, 0.8)
p = tf.paragraphs[0]
r = p.add_run(); r.text = 'Everything was paper-based, slow, and invisible to the resident.'
r.font.size = Pt(20); r.font.name = FONT; r.font.bold = True; r.font.color.rgb = GREEN
footer(s, 2)

# ---------- Slide 3: The Solution ----------
s = new_slide()
add_title(s, '02', 'The Solution')
card(s, 0.9, 1.7, 5.6, 3.6, 'For Citizens', [
    'Request services & complaints online',
    'Track status in real time',
    'Read official announcements',
    'Get notified (in-app + SMS)',
])
card(s, 6.8, 1.7, 5.6, 3.6, 'For Administrators', [
    'One dashboard for all submissions',
    'Process requests & complaints',
    'Manage residents & announcements',
    'Reports + full audit trail',
])
tf = textbox(s, 0.9, 5.7, 11.5, 0.9)
p = tf.paragraphs[0]
for txt, bold, color in [
    ('Public users can also ', False, MUTED),
    ('track submissions', True, ACCENT),
    (' and read announcements ', False, MUTED),
    ('without an account', True, ACCENT),
    ('.', False, MUTED),
]:
    r = p.add_run(); r.text = txt
    r.font.size = Pt(17); r.font.name = FONT; r.font.bold = bold; r.font.color.rgb = color
footer(s, 3)

# ---------- Slide 4: Roles & Access Control ----------
s = new_slide()
add_title(s, '03', 'Roles & Access Control')
rows = [
    ('Role', 'Access', 'Enforcement'),
    ('Public (guest)', 'Landing, announcements, tracking', 'Open routes'),
    ('Citizen', '/citizen/* portal', 'Email OTP + ID verification'),
    ('Admin', '/admin/* dashboard', 'Trusted app_metadata role'),
]
table_shape = s.shapes.add_table(4, 3, Inches(0.9), Inches(1.8), Inches(11.5), Inches(2.6))
tbl = table_shape.table
tbl.columns[0].width = Inches(2.6)
tbl.columns[1].width = Inches(4.7)
tbl.columns[2].width = Inches(4.2)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tfc = cell.text_frame
        tfc.margin_left = Inches(0.15)
        pc = tfc.paragraphs[0]
        rc = pc.add_run(); rc.text = val
        rc.font.size = Pt(16); rc.font.name = FONT
        if ri == 0:
            rc.font.bold = True; rc.font.color.rgb = ACCENT
        elif ci == 0:
            rc.font.bold = True; rc.font.color.rgb = GREEN
        else:
            rc.font.color.rgb = TEXT
tf = textbox(s, 0.9, 4.8, 11.5, 1.7)
p = tf.paragraphs[0]
rich(p, 'Two layers: **middleware redirects** at the HTTP layer + **Row-Level Security** inside the database — even bypassing the UI, a citizen can only query their own rows.', 18)

# ---------- Slide 5: Citizen Features ----------
s = new_slide()
add_title(s, '04', 'Citizen Features')
tf = textbox(s, 0.9, 1.7, 5.7, 5.0)
left = [
    '**Dashboard** — statistics & recent activity',
    '**Request a Service** — water, waste, street, health, education',
    '**File a Complaint** — evidence upload + two-way messaging',
    '**My Requests / Complaints** — status tracking with SLA',
    '**ID Verification** — upload government ID, OCR processing',
]
for i, it in enumerate(left):
    bullet(tf, it, size=16, first=(i == 0), space_after=14)
tf = textbox(s, 6.9, 1.7, 5.5, 5.0)
right = [
    '**Proxy Filing** — authorize another resident to file for you',
    '**Announcements** — official LGU updates',
    '**Document Pickups** — request & schedule release',
    '**Notifications** — in-app + SMS via Semaphore',
    '**Feedback & Settings** — language, preferences',
]
for i, it in enumerate(right):
    bullet(tf, it, size=16, first=(i == 0), space_after=14)
footer(s, 5)

# ---------- Slide 6: Admin Features ----------
s = new_slide()
add_title(s, '05', 'Admin Features')
tf = textbox(s, 0.9, 1.7, 5.7, 5.0)
left = [
    '**Dashboard & Analytics** — Leaflet heatmap of the barangay',
    '**Requests / Complaints** — process via allowed transitions',
    '**Verification Queue** — approve/reject resident IDs',
    '**Residents Directory** — residents, pre-registered, purok manifest',
]
for i, it in enumerate(left):
    bullet(tf, it, size=17, first=(i == 0), space_after=16)
tf = textbox(s, 6.9, 1.7, 5.5, 5.0)
right = [
    '**Announcements** — publish with images',
    '**Officials & Designations**',
    '**Reports** — Excel export via ExcelJS',
    '**Audit Logs** — every admin action recorded',
]
for i, it in enumerate(right):
    bullet(tf, it, size=17, first=(i == 0), space_after=16)
footer(s, 6)

# ---------- Slide 7: How a Request Flows ----------
s = new_slide()
add_title(s, '06', 'How a Request Flows')
row1 = ['Sign up', 'Email OTP', 'Verify ID (OCR)', 'Submit Request']
row2 = ['Supabase + RLS', 'Admin Processes', 'Notify (SMS)', 'Resolved']
pw, gap, x0, y0, ph = 2.6, 0.45, 0.9, 2.1, 0.75
for ri, row in enumerate([row1, row2]):
    x = x0
    y = y0 + ri * 1.4
    for ci, label in enumerate(row):
        pill(s, x, y, pw, ph, label, terminal=(ri == 1 and ci == 3))
        x += pw
        if ci < len(row) - 1:
            arrow(s, x, y, w=gap, h=ph)
            x += gap
tf = textbox(s, 0.9, 5.3, 11.5, 1.2)
p = tf.paragraphs[0]
rich(p, 'Anyone can track progress publicly at **/track** with a reference number — **no login required**.', 18)
footer(s, 7)


# ---------- Slide 8: Architecture ----------
s = new_slide()
add_title(s, '07', 'Architecture')
card(s, 0.9, 1.6, 11.5, 1.15, '', [], size=14)
tf = textbox(s, 1.2, 1.75, 11.0, 0.9)
p = tf.paragraphs[0]
rich(p, '**Client** — Browser: React 19, Tailwind, shadcn/ui', 17)
tf = textbox(s, 0.9, 2.8, 11.5, 0.5)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
rich(p, '↕ HTTP requests (middleware: auth · CSRF · rate limit)', 15, color=MUTED)
card(s, 0.9, 3.3, 11.5, 1.15, '', [], size=14)
tf = textbox(s, 1.2, 3.45, 11.0, 0.9)
p = tf.paragraphs[0]
rich(p, '**Server** — Next.js 16 App Router: API routes, server actions', 17)
tf = textbox(s, 0.9, 4.45, 11.5, 0.5)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
rich(p, '↕ Parameterized queries with Row-Level Security', 15, color=MUTED)
card(s, 0.9, 4.95, 11.5, 1.15, '', [], size=14)
tf = textbox(s, 1.2, 5.1, 11.0, 0.9)
p = tf.paragraphs[0]
rich(p, '**Data** — Supabase: Postgres + Auth + Storage', 17)
tf = textbox(s, 0.9, 6.25, 11.5, 0.5)
p = tf.paragraphs[0]
rich(p, 'External services: **Semaphore** (SMS) · **Vercel** (hosting + analytics)', 15, color=MUTED)
footer(s, 8)

# ---------- Slide 9: Tech Stack ----------
s = new_slide()
add_title(s, '08', 'Tech Stack')
card(s, 0.9, 1.7, 5.6, 4.4, 'Frontend', [
    'Next.js 16 (App Router)',
    'React 19 + TypeScript',
    'Tailwind CSS v4 + shadcn/ui',
    'React Hook Form + Zod validation',
    'Leaflet maps · TinyMCE editor',
])
card(s, 6.8, 1.7, 5.6, 4.4, 'Backend & Ops', [
    'Supabase — Postgres, Auth, Storage',
    'Row-Level Security policies',
    'Semaphore SMS (optional)',
    'Vercel deployment (GitHub CI)',
    'Playwright + Node test runner',
])
footer(s, 9)

# ---------- Slide 10: Finite State Machine ----------
s = new_slide()
add_title(s, '09', 'CS Concept: Finite State Machine')
tf = textbox(s, 0.9, 1.6, 11.5, 1.2)
p = tf.paragraphs[0]
rich(p, 'lib/status-machine.ts — lifecycles modeled as **deterministic finite automata**: states, initial state, terminal states, transition relation.', 17)
states = [('pending', False), ('processing', False), ('approved', True), ('rejected', True)]
pw, gap, x0, y, ph = 2.3, 0.7, 1.1, 3.1, 0.8
x = x0
for i, (label, term) in enumerate(states):
    pill(s, x, y, pw, ph, label, terminal=term, size=17)
    x += pw
    if i < len(states) - 1:
        arrow(s, x, y, w=gap, h=ph, txt='──▶', size=16)
        x += gap
tf = textbox(s, 0.9, 4.4, 11.5, 2.0)
p = tf.paragraphs[0]
rich(p, 'The UI only offers **legal transitions**; the server rejects illegal ones (e.g. **approved → pending**). One authoritative guard for server actions, API routes, and UI. Same pattern for **complaints** and **feedback**.', 17)
footer(s, 10)

# ---------- Slide 11: Algorithms & Data Structures ----------
s = new_slide()
add_title(s, '10', 'CS Concept: Algorithms & Data Structures')
card(s, 0.9, 1.7, 5.6, 4.7, 'Algorithms', [
    '**Levenshtein distance** (DP, O(n×m)) — fuzzy name matching in ID verification',
    '**NYSIIS phonetic matching** — name variations',
    '**Weighted scoring** — auto-verify vs. needs-review',
    '**Rate limiting** — fixed-window counters (Postgres-durable)',
], size=15)
card(s, 6.8, 1.7, 5.6, 4.7, 'Data Structures', [
    '**HashMap** — O(1) rate-limit counters, resident lookups',
    '**HashSet** — XSS allowlist, allowed transitions',
    '**Job queue** — Postgres-backed OCR jobs + polling',
    '**TTL cleanup** — DB trigger auto-deletes expired PII rows',
], size=15)
footer(s, 11)


# ---------- Slide 12: Security Layers ----------
s = new_slide()
add_title(s, '11', 'CS Concept: Security Layers')
tf = textbox(s, 0.9, 1.6, 11.5, 5.2)
items = [
    '**Cryptography** — AES-GCM encryption (random IV), SHA-256 key derivation, bcrypt hashing, CSPRNG tokens',
    '**CSRF protection** — 32-byte double-submit tokens; header must match cookie → else HTTP 403',
    '**RBAC** — roles read only from app_metadata (users cannot self-promote to admin)',
    '**Row-Level Security** — enforced inside the database, not just the UI',
    '**Rate limiting** — fixed-window counters, durable in Postgres (survives serverless cold starts)',
    '**XSS defense** — HTML tag/attribute allowlist sanitizer',
    '**PII hygiene** — DB trigger auto-deletes expired OCR rows (TTL / garbage collection)',
]
for i, it in enumerate(items):
    bullet(tf, it, size=17, first=(i == 0), space_after=13)
footer(s, 12)

# ---------- Slide 13: Testing & Quality ----------
s = new_slide()
add_title(s, '12', 'Testing & Quality')
tf = textbox(s, 0.9, 1.8, 5.7, 4.0)
left = [
    '**Unit tests** — pnpm test (Node test runner on lib/**/*.test.ts)',
    '**E2E tests** — Playwright: landing, auth guards, public tracking, citizen flows',
]
for i, it in enumerate(left):
    bullet(tf, it, size=17, first=(i == 0), space_after=16)
tf = textbox(s, 6.9, 1.8, 5.5, 4.0)
right = [
    '**Accessibility** — axe-core WCAG 2.1 AA scan; serious violations fail the build',
    '**CI/CD** — GitHub Actions on Node 22: lint, typecheck, build, test',
]
for i, it in enumerate(right):
    bullet(tf, it, size=17, first=(i == 0), space_after=16)
tf = textbox(s, 0.9, 5.6, 11.5, 0.8)
p = tf.paragraphs[0]
rich(p, 'Deployment: push to GitHub → connect to Vercel → set env vars → live.', 16, color=MUTED)
footer(s, 13)

# ---------- Slide 14: Future Work & Conclusion ----------
s = new_slide()
add_title(s, '13', 'Future Work & Conclusion')
tf = textbox(s, 0.9, 1.7, 11.5, 2.0)
items = [
    'Offline support via service worker (network-first caching already shipped)',
    'Map-based request clustering on the Leaflet heatmap',
    'Automated sync of the pre-registered resident roster',
]
for i, it in enumerate(items):
    bullet(tf, it, size=18, first=(i == 0), space_after=12)
tf = textbox(s, 0.9, 3.9, 11.5, 1.8)
p = tf.paragraphs[0]
rich(p, 'LingkodBayan turns a paper-based barangay workflow into a **transparent, secure, and accessible** digital service — built on core CS foundations: **finite state machines, dynamic programming, cryptography, and access-control theory**.', 19)
tf = textbox(s, 0.9, 5.9, 11.5, 0.8)
p = tf.paragraphs[0]
r = p.add_run(); r.text = 'Thank you!'
r.font.size = Pt(30); r.font.name = FONT; r.font.bold = True; r.font.color.rgb = GREEN
footer(s, 14)

# ---------- Embed speaker notes from SPEAKER_NOTES.md ----------
notes_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'SPEAKER_NOTES.md')
if os.path.exists(notes_path):
    md = open(notes_path, encoding='utf-8').read()
    parts = re.split(r'(?m)^## (Slide \d+) ', md)
    # parts: [preamble, 'Slide 1', body1, 'Slide 2', body2, ...]
    for idx in range(1, len(parts) - 1, 2):
        try:
            n = int(parts[idx].split()[1])
        except (ValueError, IndexError):
            continue
        if not (1 <= n <= len(prs.slides)):
            continue
        body = parts[idx + 1]
        body = re.sub(r'(?m)^## Slide \d+.*$', '', body)  # drop next header if merged
        lines = []
        for line in body.splitlines():
            line = re.sub(r'^> ?', '', line)
            line = line.replace('**', '').replace('⭐', '').strip()
            lines.append(line)
        text = '\n'.join(lines).strip()
        if text:
            prs.slides[n - 1].notes_slide.notes_text_frame.text = text

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'LingkodBayan.pptx')
prs.save(out)
print(f'Saved {out} with {len(prs.slides)} slides')

footer(s, 4)
